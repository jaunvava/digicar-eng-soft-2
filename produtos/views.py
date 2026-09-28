from django.shortcuts import render, redirect, get_object_or_404
from django.db import transaction
from decimal import Decimal, InvalidOperation
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db.models import Q
from django.utils import timezone
from django import forms
from .models import Produto, MovimentoEstoque
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


@login_required
def lista(request):
    empresa = request.empresa
    q    = request.GET.get('q', '')
    tipo = request.GET.get('tipo', '')
    situacao = request.GET.get('situacao', 'ativo')
    produtos = Produto.objects.filter(empresa=empresa)
    
    if situacao == 'ativo':
        produtos = produtos.filter(ativo=True)
    elif situacao == 'inativo':
        produtos = produtos.filter(ativo=False)

    if q:
        produtos = produtos.filter(Q(nome__icontains=q) | Q(codigo__icontains=q))
    if tipo:
        produtos = produtos.filter(tipo=tipo)

    total_registros = produtos.count()

    context = {
        'produtos': produtos,
        'q': q,
        'tipo': tipo,
        'situacao': situacao,
        'total': total_registros,
    }
    return render(request, 'produtos/lista.html', context)
    
@login_required
def imprimir(request):
    empresa = request.empresa
    q    = request.GET.get('q', '')
    tipo = request.GET.get('tipo', '')
    situacao = request.GET.get('situacao', 'ativo')
    produtos = Produto.objects.filter(empresa=empresa)
    
    if situacao == 'ativo':
        produtos = produtos.filter(ativo=True)
    elif situacao == 'inativo':
        produtos = produtos.filter(ativo=False)

    if q:
        produtos = produtos.filter(Q(nome__icontains=q) | Q(codigo__icontains=q))
    if tipo:
        produtos = produtos.filter(tipo=tipo)
        
    context = {
        'empresa':  empresa,
        'produtos': produtos,
        'total':    produtos.count(),
        'q':        q,
        'tipo':     tipo,
        'situacao': situacao,
        'now':      timezone.localtime(timezone.now()),
    }
    return render(request, 'produtos/imprimir.html', context)

from django.http import HttpResponse

@login_required
def exportar_excel(request):
    empresa = request.empresa
    q    = request.GET.get('q', '')
    tipo = request.GET.get('tipo', '')
    situacao = request.GET.get('situacao', 'ativo')
    produtos = Produto.objects.filter(empresa=empresa)
    
    if situacao == 'ativo':
        produtos = produtos.filter(ativo=True)
    elif situacao == 'inativo':
        produtos = produtos.filter(ativo=False)

    if q:
        produtos = produtos.filter(Q(nome__icontains=q) | Q(codigo__icontains=q))
    if tipo:
        produtos = produtos.filter(tipo=tipo)

    # ── Workbook ─────────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Produtos'

    COR_HEADER_FILL = '6366f1'
    COR_ROW_PAR     = 'f8f7ff'
    COR_ROW_IMPAR   = 'ffffff'
    COR_BORDA       = 'e2e2e2'

    thin = Side(style='thin', color=COR_BORDA)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Cabeçalho da empresa (linha 1-2)
    ws.merge_cells('A1:G1')
    cell_empresa = ws['A1']
    cell_empresa.value = f"{empresa.nome}  |  CNPJ: {empresa.cnpj or '—'}  |  Tel: {empresa.telefone or '—'}  |  {empresa.email or ''}"
    cell_empresa.font      = Font(bold=True, size=11, color='1e293b')
    cell_empresa.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 22

    ws.merge_cells('A2:G2')
    cell_rel = ws['A2']
    cell_rel.value = f"Relatório de Produtos  —  Gerado em: {timezone.localtime(timezone.now()).strftime('%d/%m/%Y %H:%M:%S')}"
    cell_rel.font      = Font(italic=True, size=9, color='64748b')
    cell_rel.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 16

    ws.row_dimensions[3].height = 6

    # ── Cabeçalhos das colunas (linha 4)
    headers = ['Código', 'Nome do Produto/Serviço', 'Tipo', 'Preço Custo', 'Preço Venda', 'Estoque']
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font      = Font(bold=True, color='ffffff', size=10)
        cell.fill      = PatternFill('solid', fgColor=COR_HEADER_FILL)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border    = border
    ws.row_dimensions[4].height = 20

    # ── Dados ─────────────────────────────────────────────────────────────────
    for row_idx, p in enumerate(produtos, start=5):
        fill_color = COR_ROW_PAR if row_idx % 2 == 0 else COR_ROW_IMPAR
        row_fill   = PatternFill('solid', fgColor=fill_color)

        dados_linha = [
            p.codigo or '-',
            p.nome,
            p.get_tipo_display(),
            float(p.preco_custo),
            float(p.preco_venda),
            float(p.estoque_atual) if p.tipo == 'produto' else '-'
        ]
        
        for col_idx, valor in enumerate(dados_linha, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=valor)
            cell.fill = row_fill
            cell.border = border
            cell.alignment = Alignment(vertical='center')
            if col_idx in (5, 6):
                cell.number_format = 'R$ #,##0.00'

    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 15

    ws.freeze_panes = 'A5'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="produtos.xlsx"'
    wb.save(response)
    return response


def _salvar_imagem(produto, request):
    """Grava a imagem enviada no storage configurado e apaga a anterior."""
    if request.POST.get('remover_imagem') and produto.imagem:
        produto.imagem.delete(save=True)
    arquivo = request.FILES.get('imagem')
    if not arquivo:
        return
    forms.ImageField().clean(arquivo)  # valida que o arquivo é uma imagem (Pillow)
    antiga = produto.imagem.name if produto.imagem else None
    produto.imagem.save(arquivo.name, arquivo, save=True)
    if antiga:
        produto.imagem.storage.delete(antiga)


@login_required
def novo(request):
    empresa = request.empresa
    if request.method == 'POST':
        try:
            with transaction.atomic():
                produto = Produto.objects.create(
                    empresa=empresa,
                    ativo=(request.POST.get('ativo') in ['on', 'True', 'true', '1']),
                    tipo=request.POST.get('tipo', 'produto'),
                    codigo=request.POST.get('codigo', ''),
                    nome=request.POST.get('nome'),
                    unidade=request.POST.get('unidade', 'UN'),
                    preco_custo=request.POST.get('preco_custo', 0) or 0,
                    preco_venda=request.POST.get('preco_venda', 0) or 0,
                    estoque_atual=request.POST.get('estoque_atual', 0) or 0,
                    estoque_minimo=request.POST.get('estoque_minimo', 0) or 0,
                )
                produto.refresh_from_db()  # converte os valores do POST (str) em Decimal
                if produto.tipo == 'produto' and produto.estoque_atual > 0:
                    MovimentoEstoque.objects.create(produto=produto, tipo='ajuste', quantidade=produto.estoque_atual,
                        estoque_anterior=0, estoque_posterior=produto.estoque_atual, motivo='Saldo inicial', usuario=request.user)
                # Depois do create: o caminho da imagem usa o ID do produto.
                _salvar_imagem(produto, request)
            messages.success(request, 'Produto cadastrado com sucesso!')
            return redirect('produtos:lista')
        except Exception as e:
            messages.error(request, f'Erro: {e}')
    return render(request, 'produtos/form.html', {'titulo': 'Novo Produto'})


@login_required
def editar(request, pk):
    produto = get_object_or_404(Produto, pk=pk, empresa=request.empresa)
    if request.method == 'POST':
        try:
            with transaction.atomic():
                estoque_anterior = produto.estoque_atual
                produto.ativo         = (request.POST.get('ativo') in ['on', 'True', 'true', '1'])
                produto.tipo          = request.POST.get('tipo', 'produto')
                produto.codigo        = request.POST.get('codigo', '')
                produto.nome          = request.POST.get('nome')
                produto.unidade       = request.POST.get('unidade', 'UN')
                produto.preco_custo   = request.POST.get('preco_custo', 0) or 0
                produto.preco_venda   = request.POST.get('preco_venda', 0) or 0
                produto.estoque_atual = request.POST.get('estoque_atual', 0) or 0
                produto.estoque_minimo = request.POST.get('estoque_minimo', 0) or 0
                produto.save()
                if produto.tipo == 'produto' and produto.estoque_atual != estoque_anterior:
                    MovimentoEstoque.objects.create(produto=produto, tipo='ajuste',
                        quantidade=produto.estoque_atual, estoque_anterior=estoque_anterior,
                        estoque_posterior=produto.estoque_atual, motivo='Ajuste no cadastro', usuario=request.user)
                _salvar_imagem(produto, request)
            messages.success(request, 'Produto atualizado!')
            return redirect('produtos:lista')
        except Exception as e:
            messages.error(request, f'Erro: {e}')
    return render(request, 'produtos/form.html', {'titulo': 'Editar Produto', 'produto': produto})


@login_required
def excluir(request, pk):
    produto = get_object_or_404(Produto, pk=pk, empresa=request.empresa)
    if produto.imagem:
        produto.imagem.delete(save=False)
    produto.delete()
    messages.success(request, 'Produto excluído!')
    return redirect('produtos:lista')


@login_required
def estoque(request, pk):
    produto = get_object_or_404(Produto, pk=pk, empresa=request.empresa, tipo='produto')
    if request.method == 'POST':
        try:
            tipo = request.POST.get('tipo')
            quantidade = Decimal(request.POST.get('quantidade', '0'))
            if tipo not in dict(MovimentoEstoque.TIPO_CHOICES) or quantidade <= 0:
                raise ValueError('Informe um tipo e uma quantidade positiva.')
            with transaction.atomic():
                produto = Produto.objects.select_for_update().get(pk=produto.pk)
                anterior = produto.estoque_atual
                posterior = (anterior + quantidade if tipo == 'entrada' else
                             anterior - quantidade if tipo == 'saida' else quantidade)
                if posterior < 0:
                    raise ValueError('A saída não pode deixar o estoque negativo.')
                produto.estoque_atual = posterior
                produto.save(update_fields=['estoque_atual', 'atualizado_em'])
                MovimentoEstoque.objects.create(produto=produto, tipo=tipo, quantidade=quantidade,
                    estoque_anterior=anterior, estoque_posterior=posterior,
                    motivo=request.POST.get('motivo', '').strip(), usuario=request.user)
            messages.success(request, 'Movimentação registrada com sucesso!')
            return redirect('produtos:estoque', pk=pk)
        except (InvalidOperation, ValueError) as exc:
            messages.error(request, f'Não foi possível registrar: {exc}')
    return render(request, 'produtos/estoque.html', {'produto': produto, 'movimentos': produto.movimentos.select_related('usuario')[:100]})
